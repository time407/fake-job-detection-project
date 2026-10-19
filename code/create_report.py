from __future__ import annotations

import json
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "outputs"
REPORT = ROOT / "Fake_Job_Detection_Report.docx"
NAVY = "183B5B"
PALE = "EAF1F6"
LIGHT = "F5F7F9"
GRAY = "D9D9D9"
CN_FONT = "STSong"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v)); node.set(qn("w:type"), "dxa")


def set_repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_border(cell, color=GRAY, size="5") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        node = borders.find(qn(tag))
        if node is None:
            node = OxmlElement(tag); borders.append(node)
        node.set(qn("w:val"), "single"); node.set(qn("w:sz"), size); node.set(qn("w:color"), color)


def format_run(run, size=None, bold=None, color=None, font=CN_FONT) -> None:
    run.font.name = font
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), CN_FONT)
    if size: run.font.size = Pt(size)
    if bold is not None: run.bold = bold
    if color: run.font.color.rgb = RGBColor.from_string(color)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 "); format_run(run, 9, color="666666")
    fld = OxmlElement("w:fldSimple"); fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)
    run = paragraph.add_run(" 页"); format_run(run, 9, color="666666")


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_repeat_header(table.rows[0])
    for j, h in enumerate(headers):
        cell = table.rows[0].cells[j]
        cell.text = str(h); set_cell_shading(cell, NAVY); set_cell_margins(cell); set_cell_border(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs: format_run(r, 9, True, "FFFFFF")
        if widths: cell.width = Inches(widths[j])
    for i, row in enumerate(rows):
        cells = table.add_row().cells
        for j, value in enumerate(row):
            cells[j].text = str(value)
            set_cell_margins(cells[j]); set_cell_border(cells[j])
            if i % 2: set_cell_shading(cells[j], LIGHT)
            cells[j].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cells[j].paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs: format_run(r, 9)
            if widths: cells[j].width = Inches(widths[j])
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text); format_run(r, 10.5)
    return p


def pct(v): return f"{100*v:.1f}%"


def main():
    payload = json.loads((OUT / "metrics.json").read_text(encoding="utf-8"))
    meta, models = payload["metadata"], payload["models"]
    final = models["fused"]; selected = final["test_selected"]; default = final["test_at_0_5"]
    errors = pd.read_csv(OUT / "test_errors.csv")

    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Inches(8.5); sec.page_height = Inches(11)
    sec.top_margin = Inches(0.72); sec.bottom_margin = Inches(0.68)
    sec.left_margin = Inches(0.78); sec.right_margin = Inches(0.78)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = CN_FONT; normal._element.rPr.rFonts.set(qn("w:eastAsia"), CN_FONT)
    normal.font.size = Pt(10.5); normal.font.color.rgb = RGBColor(30, 30, 30)
    normal.paragraph_format.space_after = Pt(6); normal.paragraph_format.line_spacing = 1.15
    for name, size, before, after in [("Title", 26, 0, 14), ("Heading 1", 17, 16, 8), ("Heading 2", 12.5, 11, 5)]:
        st = styles[name]; st.font.name = CN_FONT
        st._element.rPr.rFonts.set(qn("w:eastAsia"), CN_FONT)
        st.font.size = Pt(size); st.font.bold = True; st.font.color.rgb = RGBColor(0, 0, 0)
        st.paragraph_format.space_before = Pt(before); st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True
    styles["Title"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title_ppr = styles["Title"]._element.get_or_add_pPr()
    border = title_ppr.find(qn("w:pBdr"))
    if border is not None:
        title_ppr.remove(border)
    for section in doc.sections:
        add_page_number(section.footer.paragraphs[0])

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(52)
    r = p.add_run("机器学习项目报告"); format_run(r, 11, True, NAVY)
    title = doc.add_paragraph(style="Title")
    title.add_run("虚假招聘信息检测")
    sub = doc.add_paragraph()
    sub.paragraph_format.space_after = Pt(18)
    r = sub.add_run("融合文本 内容结构 缺失模式与职位属性的风险识别方法"); format_run(r, 14, False, "333333")
    doc.add_picture(str(OUT / "figures" / "precision_recall.png"), width=Inches(6.35))
    cap = doc.paragraphs[-1]; cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    info = doc.add_paragraph()
    info.paragraph_format.space_before = Pt(18)
    r = info.add_run("数据规模  17,880 条职位信息\n欺诈比例  4.84%\n最终模型  TF IDF 与结构化特征融合的加权逻辑回归\n报告日期  2026 年 10 月 5 日")
    format_run(r, 10.5, color="444444")

    doc.add_page_break()
    doc.add_heading("摘要与结论", level=1)
    p = doc.add_paragraph()
    p.add_run("建议采用融合模型作为第一阶段筛查器。 ").bold = True
    p.add_run(
        f"该模型在按重复文本分组的独立测试集上取得 ROC AUC {final['test_selected']['roc_auc']:.3f}、"
        f"PR AUC {final['test_selected']['pr_auc']:.3f}。使用验证集选出的阈值 {final['validation_selected_threshold']:.3f} 后，"
        f"召回率为 {pct(selected['recall'])}，精确率为 {pct(selected['precision'])}；在 3,450 条测试记录中，"
        f"识别出 176 条欺诈信息，漏判 11 条，同时产生 98 条误报。"
    )
    add_bullet(doc, "模型利用职位标题、公司介绍、职位描述、任职要求和福利等文本，并保留各文本段落的来源标记。")
    add_bullet(doc, "字符级特征补充拼写、网址、格式和短语变体信息；结构化特征包含类别、地点、薪资格式、布尔字段、缺失模式与文本形态。")
    add_bullet(doc, "重复招聘模板被作为同一组切分，避免相同文本同时出现在训练集与测试集，从而降低评估泄漏。")
    add_bullet(doc, "推荐把模型分数用于排序和人工复核。上线前应根据漏判成本与审核容量重新设定阈值，并监控数据漂移。")
    add_table(doc, ["指标", "结果", "解释"], [
        ["PR AUC", f"{selected['pr_auc']:.3f}", "适合衡量少数类检测；随机基线约为 0.048"],
        ["ROC AUC", f"{selected['roc_auc']:.3f}", "衡量欺诈样本整体排在正常样本之前的能力"],
        ["召回率", pct(selected["recall"]), "187 条欺诈样本中检出 176 条"],
        ["精确率", pct(selected["precision"]), "被标记样本中约三分之二确为欺诈"],
        ["平衡准确率", pct(selected["balanced_accuracy"]), "同时考虑正常类与欺诈类识别率"],
    ], widths=[1.25, 1.25, 4.1])

    doc.add_heading("数据理解与质量", level=1)
    doc.add_paragraph(
        f"数据包含 {meta['rows']:,} 条招聘信息，其中 {meta['positive_rows']} 条标记为欺诈，正类比例为 {pct(meta['positive_rate'])}。"
        "该比例意味着单纯预测所有职位均为正常也能得到约 95% 的准确率，因此准确率不适合作为主指标。"
    )
    doc.add_picture(str(OUT / "figures" / "missingness.png"), width=Inches(6.35))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph("图 1  各字段缺失比例")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.style = styles["Caption"]
    doc.add_paragraph(
        "薪资范围缺失 84.0%，部门缺失 64.6%，学历要求缺失 45.3%，福利缺失 40.3%。这些缺失并非简单删除："
        "模型使用缺失类别和缺失指示变量，因为招聘信息是否完整本身可能反映发布主体的行为。job_id 仅作记录标识，未进入模型。"
    )
    add_table(doc, ["字段组", "使用方式", "信息目的"], [
        ["五段文本", "词级一至二元 TF IDF", "语义词汇与短语"],
        ["五段文本", "字符级三至五元 TF IDF", "拼写、格式、网址及词形变体"],
        ["类别字段", "低频约束的一位有效编码", "行业、职能、经验、学历与雇佣类型"],
        ["地点", "拆分国家 地区 城市", "保留不同地理层级"],
        ["薪资", "下限 上限 区间宽度 缺失", "提取数值和填写行为"],
        ["文本形态", "长度 词数 大写 数字 邮件 网址 感叹号", "补充内容风格信号"],
        ["布尔字段", "直接输入", "远程 公司标志 筛选问题"],
    ], widths=[1.2, 2.25, 3.15])

    doc.add_heading("验证设计", level=1)
    doc.add_paragraph(
        f"五段核心文本标准化后形成 {meta['unique_text_groups']:,} 个唯一组。共有约 2,131 条记录与其他记录共享完整文本。"
        "分组切分确保同一文本组只属于训练、验证或测试中的一个集合。训练集拟合特征和模型，验证集选择决策阈值，测试集只用于最终报告。"
    )
    add_table(doc, ["数据集", "记录数", "欺诈比例", "用途"], [
        ["训练集", f"{meta['split_sizes']['train']:,}", pct(meta['split_positive_rates']['train']), "拟合向量器与分类器"],
        ["验证集", f"{meta['split_sizes']['validation']:,}", pct(meta['split_positive_rates']['validation']), "选择最大化 F2 的阈值"],
        ["测试集", f"{meta['split_sizes']['test']:,}", pct(meta['split_positive_rates']['test']), "一次性最终评估"],
    ], widths=[1.2, 1.2, 1.2, 3.0])
    doc.add_paragraph(
        "F2 分数对召回率的权重高于精确率，符合初筛系统中漏掉欺诈信息通常比增加一次人工复核更昂贵的假设。"
        "若实际业务成本不同，应更换阈值目标，而不必重新训练全部文本表示。"
    )

    doc.add_heading("模型设计", level=1)
    doc.add_heading("特征融合", level=2)
    doc.add_paragraph(
        "词级向量捕捉招聘语义，字符级向量对拼写错误、网址占位符、词形和标点模式更敏感。结构化分支通过一位有效编码和标准化数值变量提供文本之外的证据。"
        "三类稀疏向量拼接后输入带类别权重的逻辑回归，共形成 201,268 个特征。"
    )
    doc.add_heading("分类器与可解释性", level=2)
    doc.add_paragraph(
        "逻辑回归适合高维稀疏特征，训练稳定且能输出连续风险分数。类别权重用于缓解正负样本不平衡。"
        "系数可用于定位与风险分数相关的词语、字符片段和职位属性，但系数表示数据内的相关性，不能解释为因果关系，也不应把地点或行业标签单独当作欺诈证据。"
    )
    doc.add_heading("决策阈值", level=2)
    doc.add_paragraph(
        f"验证集上最大化 F2 的阈值为 {final['validation_selected_threshold']:.3f}。模型分数高于该阈值即进入高风险队列。"
        "阈值降低会提高召回率并增加误报；阈值提高则减少审核量，但会漏掉更多欺诈信息。"
    )

    doc.add_heading("实验结果", level=1)
    doc.add_picture(str(OUT / "figures" / "model_comparison.png"), width=Inches(6.25))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph("图 2  分组测试集上的模型比较")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.style = styles["Caption"]
    rows = []
    labels = {"structured": "仅结构化特征", "word_text": "仅词级文本", "fused": "融合模型"}
    for key in ("structured", "word_text", "fused"):
        m = models[key]["test_selected"]
        rows.append([labels[key], f"{m['pr_auc']:.3f}", f"{m['roc_auc']:.3f}", pct(m['precision']), pct(m['recall']), f"{m['f2']:.3f}"])
    add_table(doc, ["模型", "PR AUC", "ROC AUC", "精确率", "召回率", "F2"], rows, widths=[1.55, 0.9, 0.9, 0.9, 0.9, 0.8])
    ci = final["bootstrap_95_ci"]
    doc.add_paragraph(
        f"融合模型的测试 PR AUC 为 {selected['pr_auc']:.3f}，自助法 95% 区间为 {ci['pr_auc'][0]:.3f} 至 {ci['pr_auc'][1]:.3f}；"
        f"ROC AUC 的 95% 区间为 {ci['roc_auc'][0]:.3f} 至 {ci['roc_auc'][1]:.3f}。"
        "词级文本模型的 PR AUC 略高，但融合模型在选定阈值下得到更高召回率，且结构化分支为文本较短或缺失的职位提供补充证据。"
    )
    add_table(doc, ["阈值方案", "精确率", "召回率", "误报", "漏判", "适用情形"], [
        ["0.299 验证选择", pct(selected["precision"]), pct(selected["recall"]), "98", "11", "优先降低漏判"],
        ["0.500 默认", pct(default["precision"]), pct(default["recall"]), "40", "16", "审核资源较紧"],
    ], widths=[1.4, 0.9, 0.9, 0.7, 0.7, 2.0])

    doc.add_heading("错误分析", level=1)
    doc.add_paragraph(
        "选定阈值下共有 109 个错误，其中 11 个漏判、98 个误报。漏判通常具有较中性的职位名称或较少可疑文本模式；"
        "误报则可能与训练数据中的高风险模板、地点、行业或格式相似。以下记录用于定向复核，不代表这些职位名称或地区本身可疑。"
    )
    fn = errors[errors.error_type == "false_negative"].sort_values("score", ascending=False).head(5)
    fp = errors[errors.error_type == "false_positive"].sort_values("score", ascending=False).head(5)
    add_table(doc, ["类型", "job id", "职位名称", "风险分数"],
              [["漏判", int(r.job_id), str(r.title)[:46], f"{r.score:.3f}"] for _, r in fn.iterrows()] +
              [["误报", int(r.job_id), str(r.title)[:46], f"{r.score:.3f}"] for _, r in fp.iterrows()],
              widths=[0.8, 0.8, 4.2, 0.9])
    doc.add_paragraph(
        "下一轮改进可针对漏判样本增加难例权重，或在新增标注数据后训练更强的语义编码器。任何改进都应沿用分组测试原则，并保留当前测试集作为不可调参的基准。"
    )

    doc.add_heading("落地方案", level=1)
    add_bullet(doc, "输入校验  检查字段、数据类型、文本编码和缺失率；保留原始 job_id 便于追踪。")
    add_bullet(doc, "批量评分  输出 0 至 1 的风险分数，并使用经过业务确认的阈值生成复核队列。")
    add_bullet(doc, "人工处置  高风险记录先进入人工审核，不建议直接自动删除或拒绝发布。")
    add_bullet(doc, "反馈闭环  保存审核结论、申诉结果和最终处置，用于周期性重新训练与阈值校准。")
    add_bullet(doc, "运行监控  每周监控风险分数分布、报警率、字段缺失率与人工确认率；有标签后监控精确率和召回率。")
    add_bullet(doc, "公平检查  分地区、行业和雇佣类型评估误报率，防止历史数据中的集中模式变成不合理的自动规则。")
    doc.add_paragraph(
        "项目已包含训练脚本、预测脚本、固定切分索引、测试错误表、特征系数、评估图和序列化模型。"
        "新数据可通过 predict.py 生成 fraud_risk_score 和 fraud_prediction 两列。"
    )

    doc.add_heading("限制与后续工作", level=1)
    add_bullet(doc, "标签来自现有数据集，可能包含标注错误，也可能不能代表当前平台上的新型诈骗。")
    add_bullet(doc, "测试集与训练集来自同一数据来源；跨平台、跨时间和跨语言泛化仍需独立验证。")
    add_bullet(doc, "模型分数尚未做概率校准，适合风险排序，不应直接解释为真实欺诈概率。")
    add_bullet(doc, "高维文本系数可解释关联，但无法证明某个词、地区或行业导致欺诈。")
    add_bullet(doc, "重复文本分组只处理完全标准化后的重复；近似改写模板仍可能跨集合出现。后续可用语义相似度聚类进行更严格的去重评估。")
    doc.add_paragraph(
        "建议下一阶段收集按时间排序的新数据，建立时间外测试集，并比较校准后的线性模型、梯度提升结构化模型和预训练语言模型。"
        "只有在外部测试、审核成本和公平性指标都满足要求后，才应扩大自动化范围。"
    )

    doc.add_heading("复现说明", level=1)
    add_table(doc, ["文件", "作用"], [
        ["code/train.py", "特征工程 分组切分 训练 评估 保存模型和图表"],
        ["code/predict.py", "对新 CSV 批量评分"],
        ["outputs metrics.json", "完整验证与测试指标"],
        ["outputs split_indices.json", "固定训练 验证 测试索引"],
        ["outputs test_errors.csv", "测试集误报与漏判记录"],
        ["outputs feature_importance.csv", "最终模型正负方向系数"],
        ["outputs fake_job_detector.joblib", "向量器 结构化转换器 分类器与阈值"],
    ], widths=[2.3, 4.4])
    doc.add_paragraph(
        "复现命令为 python code/train.py --data data/fake_job_postings.csv --output outputs。"
        "随机种子固定为 42，切分索引随结果一并保存。软件依赖记录在 requirements.txt。"
    )

    doc.core_properties.title = "虚假招聘信息检测机器学习项目报告"
    doc.core_properties.subject = "Fake job detection model design and evaluation"
    doc.core_properties.author = ""
    doc.save(REPORT)
    print(REPORT)


if __name__ == "__main__":
    main()
